import sys
import os

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    DARK_BG = RGBColor(5, 5, 10)
    ORANGE = RGBColor(255, 107, 0)
    AMBER = RGBColor(255, 140, 0)
    CYAN = RGBColor(0, 200, 255)
    WHITE = RGBColor(240, 240, 245)
    MUTED = RGBColor(160, 160, 192)

    blank_slide_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text):
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.8))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.bold = True
        p.font.size = Pt(28)
        p.font.color.rgb = AMBER
        p.font.name = 'Arial'

        txBox2 = slide.shapes.add_textbox(Inches(9.5), Inches(0.4), Inches(3.0), Inches(0.5))
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = category_text
        p2.font.size = Pt(12)
        p2.font.color.rgb = CYAN
        p2.alignment = PP_ALIGN.RIGHT

    # SLIDE 1: Title
    slide1 = prs.slides.add_slide(blank_slide_layout)
    txBox = slide1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(4.5))
    tf = txBox.text_frame
    
    p = tf.paragraphs[0]
    p.text = "SURYADRISHTI"
    p.font.bold = True
    p.font.size = Pt(44)
    p.font.color.rgb = AMBER
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "Aditya-L1 Solar Flare Intelligence & Space Weather Nowcasting Center"
    p2.font.size = Pt(22)
    p2.font.color.rgb = WHITE
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf.add_paragraph()
    p3.text = "A Sub-Second Physics-Informed AI (PINN) Defense System for Satellite Avionics"
    p3.font.size = Pt(16)
    p3.font.color.rgb = CYAN
    p3.alignment = PP_ALIGN.CENTER

    p4 = tf.add_paragraph()
    p4.text = "\nPresenter: Deep Shekhar Halder (Solo Innovation) | India AI Impact Festival 2026"
    p4.font.size = Pt(14)
    p4.font.color.rgb = MUTED
    p4.alignment = PP_ALIGN.CENTER

    # SLIDE 2: Problem
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide2, "02. Critical Threat: High-Energy Space Weather", "PROBLEM STATEMENT")

    # SLIDE 3: Innovation
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide3, "03. The Breakthrough: Physics-Informed Neural AI", "INNOVATION & PHYSICS")

    # SLIDE 4: Architecture
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide4, "04. End-to-End Pipeline Architecture", "SYSTEM ARCHITECTURE")

    # SLIDE 5: 3D Globe
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide5, "05. Interactive 3D WebGL Globe Dashboard", "3D MISSION CONTROL")

    # SLIDE 6: 2D Dashboard
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide6, "06. 2D Streamlit Scientific Analytics Center", "2D ANALYTICS HUD")

    # SLIDE 7: Benchmarks
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide7, "07. Model Performance & Benchmark Results", "BENCHMARKS & METRICS")

    # SLIDE 8: Impact
    slide8 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide8, "08. Impact for AI-Powered Bharat", "NATIONAL IMPACT")

    # SLIDE 9: Q&A Shield
    slide9 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide9, "09. Pro-Defense Against Tough Judge Questions", "TECHNICAL Q&A SHIELD")

    # SLIDE 10: Conclusion
    slide10 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide10, "10. Future Roadmap & Mission Vision", "ROADMAP & CONCLUSION")

    output_path = os.path.join(os.getcwd(), "SuryaDrishti_Official_Presentation.pptx")
    prs.save(output_path)
    print(f"SUCCESSFULLY_CREATED_PPTX: {output_path}")

except Exception as e:
    print(f"PPTX_ERROR: {e}")
