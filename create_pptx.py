import sys
import subprocess

try:
    import pptx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    import pptx

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette - Schneider Electric theme
    SE_GREEN = RGBColor(0, 150, 57)
    SE_BRIGHT_GREEN = RGBColor(48, 209, 88)
    DARK_BG = RGBColor(11, 19, 31)
    CARD_BG = RGBColor(21, 34, 56)
    TEXT_LIGHT = RGBColor(248, 250, 252)
    TEXT_MUTED = RGBColor(148, 163, 184)
    WHITE = RGBColor(255, 255, 255)

    blank_slide_layout = prs.slide_layouts[6]

    slides_data = [
        {
            "title": "SuryaGrid: Physics-Informed AI for Smart Grid Resilience",
            "subtitle": "Schneider Electric Yuva Yodha Energy Tech Hackathon 2026 | Category: Grid Reliability",
            "bullets": [
                "Team Lead: Lidiya (duddekuntalidiya@gmail.com)",
                "Team Member: Deep Halder (deephalder209@gmail.com)",
                "Team Member: Sayan Pramanik (pramaniksayak145@gmail.com)",
                "Team Member: Rituraj Saha (saharituraj805@gmail.com | GitHub: RiturajtheCoder)"
            ]
        },
        {
            "title": "Problem Statement: Modern Power Grid Vulnerabilities",
            "subtitle": "The threat of extreme space weather (GMD/GIC) and renewable power instability",
            "bullets": [
                "Transformer Saturation & Destruction: Geomagnetically Induced Currents (GIC) from solar storms saturate transformer cores, triggering localized thermal runaways and rapid burnout.",
                "Legacy SCADA Invisibility: Existing SCADA infrastructure provides purely reactive metrics, lacking sub-second predictive physics modeling.",
                "Multibillion Dollar Impact: A single major solar storm or cascading blackout can cause up to $10B+ in transformer damage and weeks of power grid downtime."
            ]
        },
        {
            "title": "Alignment to Schneider Electric Grid Reliability Challenge",
            "subtitle": "Empowering Next-Generation Smart Grids with Zero-Outage Resilience",
            "bullets": [
                "Asset Performance Management: Continuous thermal and magnetic health tracking of EHV (Extra High Voltage) grid transformers.",
                "Proactive Grid Defense: Early warning system predicting GIC surges 15-45 minutes before line entry.",
                "Schneider Ecosystem Readiness: Seamless REST API architecture designed for integration with Schneider EcoStruxure platforms."
            ]
        },
        {
            "title": "Proposed Solution: SuryaGrid Physics-Informed AI",
            "subtitle": "Fusing Satellite Telemetry with Physics-Informed Neural Networks (PINNs)",
            "bullets": [
                "PINN Engine: Integrates Maxwell's Electromagnetic Equations directly into neural loss functions to eliminate AI hallucination.",
                "Advance Early Warning: Predicts localized geomagnetic flux rate (dB/dt) and transformer saturation 15-45 minutes ahead.",
                "3D Substation Twin: Real-time interactive command center for power dispatch operators and engineers."
            ]
        },
        {
            "title": "Key Features & Operator Journey",
            "subtitle": "End-to-End Workflow from Threat Sensing to Automated Neutralization",
            "bullets": [
                "1. Solar Wind Ingestion: Streams real-time solar magnetic vector (Bz) and density from ISRO Aditya-L1 / NOAA satellites.",
                "2. Physics Prediction: Computes GIC magnitude, top-oil thermal rise, and core saturation index.",
                "3. Visual Alerting: Displays live risk heatmaps on a 3D digital twin substation model.",
                "4. Automated Defense: Recommends neutral grounding resistor switching & dynamic reactive power compensation."
            ]
        },
        {
            "title": "Technical Approach & Architecture",
            "subtitle": "Robust, Latency-Optimized Architecture for High-Reliability Power Systems",
            "bullets": [
                "Data Layer: Multi-source stream processing of L1 satellite telemetry, ground magnetometers, and SCADA IoT feeds.",
                "AI/ML Core: PyTorch PINNs + LSTM + Graph Neural Networks (GNN) modeling power grid node topology.",
                "Control Interface: Fast API backend with WebSocket live streaming and 3D Three.js substation visualization."
            ]
        },
        {
            "title": "Innovation & Key Differentiators",
            "subtitle": "Why SuryaGrid Leads Over Standard Black-Box AI Models",
            "bullets": [
                "Physics-Constrained AI: Built-in physical law constraints guarantee mathematical validity during unprecedented solar storms.",
                "30-Minute Advance Warning: Moves grid management from reactive tripping to proactive surge mitigation.",
                "Aditya-L1 Integration: Pioneer solution utilizing Indian space telemetry for terrestrial energy infrastructure defense."
            ]
        },
        {
            "title": "Expected Impact & Key Beneficiaries",
            "subtitle": "Quantifiable Value Delivered to Power Utilities and Society",
            "bullets": [
                "Power Utilities (DISCOMs): Saves EHV transformers worth $5M-$10M each from catastrophic thermal burnout.",
                "Industrial Energy Consumers: Protects mission-critical facilities (data centers, hospitals, factories) from abrupt grid collapses.",
                "Sustainability Impact: Reduces reliance on high-emission diesel backup generators during blackouts."
            ]
        },
        {
            "title": "Implementation & Rollout Roadmap",
            "subtitle": "4-Phase Commercial Strategy for Grid-Scale Deployment",
            "bullets": [
                "Phase 1 (Q1-Q2): Model calibration on NOAA historical storm benchmarks & PINN fine-tuning.",
                "Phase 2 (Q3): Pilot hardware testing with localized IoT magnetometer sensors and EcoStruxure simulation.",
                "Phase 3 (Q4): Substation deployment across 5 high-vulnerability regional grid hubs.",
                "Phase 4 (Year 2): Full commercial rollout to national power transmission networks and global utilities."
            ]
        },
        {
            "title": "Team Introduction & Conclusion",
            "subtitle": "SuryaGrid Team - Ready to Build the Future of Energy Tech",
            "bullets": [
                "Lidiya (Team Lead) - duddekuntalidiya@gmail.com",
                "Deep Halder (Team Member) - deephalder209@gmail.com",
                "Sayan Pramanik (Team Member) - pramaniksayak145@gmail.com",
                "Rituraj Saha (Team Member) - saharituraj805@gmail.com",
                "Together, shaping resilient, intelligent energy grids for Schneider Electric Yuva Yodha 2026."
            ]
        }
    ]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # Background shape
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()

        # Top Accent Line
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.5), Inches(11.733), Inches(0.06))
        accent.fill.solid()
        accent.fill.fore_color.rgb = SE_GREEN
        accent.line.fill.background()

        # Slide Title
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.733), Inches(1.0))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = data["title"]
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.name = "Arial"

        # Subtitle
        p2 = tf.add_paragraph()
        p2.text = data["subtitle"]
        p2.font.size = Pt(16)
        p2.font.color.rgb = SE_BRIGHT_GREEN
        p2.font.name = "Arial"

        # Content Card / Box
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(11.733), Inches(4.9))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = SE_GREEN
        box.line.width = Pt(1.5)

        # Bullets inside Card
        tb = slide.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(11.133), Inches(4.5))
        tf_b = tb.text_frame
        tf_b.word_wrap = True
        
        for idx, bullet in enumerate(data["bullets"]):
            p_b = tf_b.paragraphs[0] if idx == 0 else tf_b.add_paragraph()
            p_b.text = "• " + bullet
            p_b.font.size = Pt(18)
            p_b.font.color.rgb = TEXT_LIGHT
            p_b.font.name = "Arial"
            p_b.space_after = Pt(14)

    output_path = "d:\\PS15_SolarFlare\\SuryaGrid_Schneider_Yuva_Yodha_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_presentation()
